import base64
import hashlib
import io
import json
import os
import time
import uuid
from itertools import zip_longest

import requests
from docx import Document
from flask import Flask, flash, redirect, render_template, request, send_file, session, url_for


LANGUAGES = {
    "zh": "中文",
    "en": "英语",
    "jp": "日语",
    "kor": "韩语",
    "fra": "法语",
    "de": "德语",
    "ru": "俄语",
    "spa": "西班牙语",
}

DEFAULT_CONFIG_PATH = os.environ.get(
    "IMAGE_TEXT_APP_CONFIG_PATH",
    os.path.join(os.path.dirname(__file__), "config.json"),
)

ENV_MAPPING = {
    "baidu_ocr_api_key": "BAIDU_OCR_API_KEY",
    "baidu_ocr_secret_key": "BAIDU_OCR_SECRET_KEY",
    "baidu_translate_app_id": "BAIDU_TRANSLATE_APP_ID",
    "baidu_translate_secret_key": "BAIDU_TRANSLATE_SECRET_KEY",
}


class AccessTokenCache:
    """Cache Baidu OCR access token to reduce authentication requests."""

    def __init__(self) -> None:
        self._token: str | None = None
        self._expires_at: float = 0.0

    def get_token(self, api_key: str, secret_key: str) -> str:
        now = time.time()
        if self._token and now < self._expires_at - 60:
            return self._token

        token_url = "https://aip.baidubce.com/oauth/2.0/token"
        resp = requests.get(
            token_url,
            params={
                "grant_type": "client_credentials",
                "client_id": api_key,
                "client_secret": secret_key,
            },
            timeout=10,
        )
        resp.raise_for_status()
        data = resp.json()
        if "access_token" not in data:
            error_msg = data.get("error_description") or "无法获取 OCR access token"
            raise RuntimeError(error_msg)

        self._token = data["access_token"]
        self._expires_at = now + float(data.get("expires_in", 0))
        return self._token


def load_config(path: str) -> dict:
    if not os.path.exists(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8") as cfg_file:
            return json.load(cfg_file)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"配置文件解析失败: {path}") from exc


CONFIG = load_config(DEFAULT_CONFIG_PATH)
token_cache = AccessTokenCache()

app = Flask(__name__, template_folder="templates", static_folder="static")
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "change-this-secret")


def get_setting(key: str) -> str | None:
    env_key = ENV_MAPPING[key]
    if env_key in os.environ:
        return os.environ[env_key]
    return CONFIG.get(key)


def get_missing_credentials() -> list[str]:
    missing = []
    required = [
        ("Baidu OCR API Key", "baidu_ocr_api_key"),
        ("Baidu OCR Secret Key", "baidu_ocr_secret_key"),
        ("Baidu 翻译 App ID", "baidu_translate_app_id"),
        ("Baidu 翻译密钥", "baidu_translate_secret_key"),
    ]
    for label, key in required:
        if not get_setting(key):
            missing.append(label)
    return missing


def call_baidu_ocr(image_bytes: bytes) -> str:
    api_key = get_setting("baidu_ocr_api_key")
    secret_key = get_setting("baidu_ocr_secret_key")
    if not api_key or not secret_key:
        raise RuntimeError("未配置 Baidu OCR 凭证")

    access_token = token_cache.get_token(api_key, secret_key)
    request_url = (
        "https://aip.baidubce.com/rest/2.0/ocr/v1/general_basic"
        f"?access_token={access_token}"
    )
    headers = {"Content-Type": "application/x-www-form-urlencoded"}
    data = {"image": base64.b64encode(image_bytes).decode("utf-8")}
    response = requests.post(request_url, data=data, headers=headers, timeout=15)
    response.raise_for_status()
    result = response.json()

    if "error_code" in result:
        message = result.get("error_msg", "未知错误")
        raise RuntimeError(f"Baidu OCR 调用失败: {message}")

    words = [item.get("words", "") for item in result.get("words_result", [])]
    return "\n".join(word for word in words if word)


def translate_text(text: str, target_language: str) -> str:
    if not text.strip():
        return ""

    app_id = get_setting("baidu_translate_app_id")
    secret_key = get_setting("baidu_translate_secret_key")
    if not app_id or not secret_key:
        raise RuntimeError("未配置 Baidu 翻译凭证")

    translate_url = "https://fanyi-api.baidu.com/api/trans/vip/translate"
    salt = uuid.uuid4().hex
    sign_raw = f"{app_id}{text}{salt}{secret_key}"
    sign = hashlib.md5(sign_raw.encode("utf-8")).hexdigest()
    payload = {
        "q": text,
        "from": "auto",
        "to": target_language,
        "appid": app_id,
        "salt": salt,
        "sign": sign,
    }
    response = requests.post(translate_url, data=payload, timeout=15)
    response.raise_for_status()
    data = response.json()

    if "error_code" in data:
        message = data.get("error_msg", "翻译接口调用失败")
        raise RuntimeError(f"Baidu 翻译失败: {message}")

    translations = [item.get("dst", "") for item in data.get("trans_result", [])]
    return "\n".join(translations)


def generate_docx(source_text: str, translated_text: str) -> io.BytesIO:
    document = Document()
    document.add_heading("图片文字识别与翻译", level=1)
    table = document.add_table(rows=1, cols=2)
    header_cells = table.rows[0].cells
    header_cells[0].text = "原文"
    header_cells[1].text = "翻译"

    for src_line, dst_line in zip_longest(
        source_text.splitlines(), translated_text.splitlines(), fillvalue=""
    ):
        row_cells = table.add_row().cells
        row_cells[0].text = src_line
        row_cells[1].text = dst_line

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def allowed_file(filename: str) -> bool:
    allowed_extensions = {"png", "jpg", "jpeg", "bmp", "gif", "tif", "tiff"}
    return "." in filename and filename.rsplit(".", 1)[1].lower() in allowed_extensions


@app.route("/", methods=["GET", "POST"])
def index():
    context = {"languages": LANGUAGES, "selected_language": "en"}
    missing_credentials = get_missing_credentials()
    context["missing_credentials"] = missing_credentials

    if request.method == "GET":
        stored_source = session.get("source_text")
        stored_translation = session.get("translated_text")
        stored_language = session.get("target_language", "en")
        if stored_source or stored_translation:
            context.update(
                {
                    "source_text": stored_source,
                    "translated_text": stored_translation,
                    "selected_language": stored_language,
                    "has_result": True,
                }
            )

    if request.method == "POST":
        if missing_credentials:
            readable = "、".join(missing_credentials)
            flash(f"请先配置以下凭证后再使用：{readable}", "error")
            return render_template("index.html", **context)

        uploaded_file = request.files.get("image")
        target_language = request.form.get("target_language", "en")

        if not uploaded_file or uploaded_file.filename == "":
            flash("请上传图片文件", "error")
            return render_template("index.html", **context)

        if not allowed_file(uploaded_file.filename):
            flash("仅支持常见图片格式 (png/jpg/jpeg/bmp/gif/tiff)", "error")
            return render_template("index.html", **context)

        image_bytes = uploaded_file.read()
        try:
            original_text = call_baidu_ocr(image_bytes)
            translation_text = translate_text(original_text, target_language)
        except Exception as exc:  # pylint: disable=broad-except
            flash(str(exc), "error")
            return render_template("index.html", **context)

        if not original_text:
            flash("未识别到任何文字", "warning")

        session["source_text"] = original_text
        session["translated_text"] = translation_text
        session["target_language"] = target_language

        context.update(
            {
                "source_text": original_text,
                "translated_text": translation_text,
                "selected_language": target_language,
                "has_result": True,
            }
        )

    return render_template("index.html", **context)


@app.route("/download/txt/<variant>", methods=["GET"])
def download_txt(variant: str):
    source_text = session.get("source_text", "")
    translated_text = session.get("translated_text", "")

    if not source_text and not translated_text:
        flash("当前没有可下载的内容，请先上传并识别图片。", "error")
        return redirect(url_for("index"))

    if variant == "source":
        filename = "ocr_text.txt"
        content = source_text
    elif variant == "translation":
        filename = "translation.txt"
        content = translated_text
    elif variant == "pair":
        filename = "ocr_translation_pair.txt"
        lines = zip_longest(
            source_text.splitlines(),
            translated_text.splitlines(),
            fillvalue="",
        )
        content = "\n".join(f"{src}\t{dst}" for src, dst in lines)
    else:
        flash("未知的下载类型", "error")
        return redirect(url_for("index"))

    buffer = io.BytesIO()
    buffer.write(content.encode("utf-8"))
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="text/plain",
    )


@app.route("/download/docx", methods=["GET"])
def download_docx():
    source_text = session.get("source_text")
    translated_text = session.get("translated_text")
    target_language = session.get("target_language", "en")

    if not source_text and not translated_text:
        flash("当前没有可下载的内容，请先上传并识别图片。", "error")
        return redirect(url_for("index"))

    buffer = generate_docx(source_text or "", translated_text or "")
    filename = f"ocr_translation_{target_language}.docx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )


if __name__ == "__main__":
    # Enable Flask debug only when explicitly requested via environment variable.
    debug_mode = os.environ.get("FLASK_DEBUG", "0") == "1"
    app.run(host="0.0.0.0", port=5000, debug=debug_mode)
